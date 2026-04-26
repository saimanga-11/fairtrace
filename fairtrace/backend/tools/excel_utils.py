from __future__ import annotations

import re
from pathlib import Path
from typing import Optional

try:
    from openpyxl import load_workbook
except Exception:  # pragma: no cover
    load_workbook = None


PREFERRED_DATA_SHEETS = {
    "original_data",
    "original data",
    "data",
    "dataset",
    "sheet1",
    "sheet",
    "source_data",
    "source data",
}

REPORT_LIKE_SHEETS = {
    "bias_report",
    "charts",
    "updated_data",
    "updated data",
    "suggested_data",
    "suggested data",
    "suggested_changes",
    "suggested changes",
}


def _normalize_sheet_name(name: str) -> str:
    text = str(name or "").strip().lower()
    text = re.sub(r"[\s_-]+", "_", text)
    return text


def select_xlsx_sheet_name(path: Path) -> str:
    if load_workbook is None:
        return "Sheet1"

    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        sheet_names = workbook.sheetnames or ["Sheet1"]

        normalized_map = {_normalize_sheet_name(name): name for name in sheet_names}

        for preferred in ("original_data", "updated_data", "suggested_data"):
            if preferred in normalized_map:
                return normalized_map[preferred]

        non_report_sheets = [
            name for name in sheet_names
            if _normalize_sheet_name(name) not in REPORT_LIKE_SHEETS
        ]
        if non_report_sheets:
            return max(
                non_report_sheets,
                key=lambda name: (workbook[name].max_row or 0) * (workbook[name].max_column or 0),
            )

        return sheet_names[0]
    finally:
        workbook.close()


def find_header_row_index(path: Path, sheet_name: Optional[str] = None) -> int:
    if load_workbook is None:
        return 0

    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        target_sheet = sheet_name or select_xlsx_sheet_name(path)
        if target_sheet not in workbook.sheetnames:
            target_sheet = workbook.sheetnames[0] if workbook.sheetnames else "Sheet1"
        worksheet = workbook[target_sheet]

        for idx, row in enumerate(worksheet.iter_rows(max_row=8, values_only=True)):
            values = [str(v).strip() if v is not None else "" for v in row]
            non_empty = [v for v in values if v]
            numeric_count = sum(1 for v in non_empty if _looks_numeric(v))
            if len(non_empty) >= 3 and numeric_count == 0:
                looks_like_header = all(len(v.split()) <= 6 for v in non_empty)
                if looks_like_header:
                    return idx
        return 0
    finally:
        workbook.close()


def _looks_numeric(value: str) -> bool:
    try:
        float(str(value).replace(",", "").strip())
        return True
    except Exception:
        return False
