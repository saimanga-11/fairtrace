import csv
from collections import Counter
from pathlib import Path
from typing import Any, Dict, List

try:
    import pandas as pd
except Exception:  # pragma: no cover
    pd = None

try:
    from openpyxl import load_workbook
except Exception:  # pragma: no cover
    load_workbook = None

from tools.excel_utils import find_header_row_index, select_xlsx_sheet_name


def _is_number(value: str) -> bool:
    try:
        float(value)
        return True
    except Exception:
        return False


def _normalize_value(value: Any) -> str:
    if value is None:
        return "missing"
    text = str(value).strip()
    return text if text else "missing"


def analyze_csv(csv_path: str) -> Dict[str, Any]:
    path = Path(csv_path)
    if path.suffix.lower() == ".xlsx":
        if load_workbook is None:
            raise RuntimeError("openpyxl is required to read .xlsx files.")
        sheet_name = select_xlsx_sheet_name(path)
        header_row = find_header_row_index(path, sheet_name)
        if pd is not None:
            dataframe = pd.read_excel(path, sheet_name=sheet_name, header=header_row)
            rows = dataframe.to_dict(orient="records")
            fieldnames = list(dataframe.columns)
        else:
            workbook = load_workbook(path, data_only=True, read_only=True)
            try:
                worksheet = workbook[sheet_name] if sheet_name in workbook.sheetnames else workbook.active
                rows_iter = worksheet.iter_rows(values_only=True)
                for _ in range(header_row):
                    next(rows_iter, None)
                headers = [
                    str(cell).strip() if cell is not None and str(cell).strip() else f"column_{index + 1}"
                    for index, cell in enumerate(next(rows_iter, []))
                ]
                rows = [
                    {
                        headers[index]: value
                        for index, value in enumerate(row)
                        if index < len(headers)
                    }
                    for row in rows_iter
                ]
                fieldnames = headers
            finally:
                workbook.close()
    elif path.suffix.lower() == ".xls":
        if pd is None:
            raise RuntimeError("Install pandas and xlrd to read .xls files, or save the file as .xlsx.")
        dataframe = pd.read_excel(path)
        rows = dataframe.to_dict(orient="records")
        fieldnames = list(dataframe.columns)
    else:
        with path.open("r", newline="", encoding="utf-8-sig") as handle:
            reader = csv.DictReader(handle)
            rows = list(reader)
        fieldnames = reader.fieldnames or []

    columns: List[Dict[str, Any]] = []
    total_rows = len(rows)

    for field in fieldnames:
        values = [_normalize_value(row.get(field)) for row in rows]
        non_missing = [value for value in values if value != "missing"]
        counts = Counter(values)
        total = sum(counts.values()) or 1
        distribution = {
            key: round((count / total) * 100, 2)
            for key, count in counts.most_common()
        }
        dominant_share = max(distribution.values()) if distribution else 0
        imbalance_flag = dominant_share > 70.0
        proxy_flag = any(token in field.lower() for token in ["zip", "postcode", "postal"])
        columns.append(
            {
                "column": field,
                "distribution": distribution,
                "imbalance_flag": imbalance_flag,
                "proxy_flag": proxy_flag,
                "dominant_share": dominant_share,
                "value_kind": "numeric" if non_missing and all(_is_number(v) for v in non_missing) else "categorical",
            }
        )

    return {
        "row_count": total_rows,
        "columns": columns,
    }
