from __future__ import annotations

import os
import math
import random
import re
import json
import tempfile
from pathlib import Path
from typing import Any, Dict, List, Tuple

try:
    import pandas as pd
except Exception:  # pragma: no cover
    pd = None

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows

try:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
except Exception:  # pragma: no cover
    plt = None

from agents.data_profiler import _clean_dataframe
from tools.excel_utils import find_header_row_index, select_xlsx_sheet_name
from tools.mistral_tool import complete_json


BACKEND_DIR = Path(__file__).resolve().parents[1]
REPORT_DIR = BACKEND_DIR / "generated_reports"
REPORT_DIR.mkdir(parents=True, exist_ok=True)

HEADER_FILL = PatternFill(fill_type="solid", fgColor="1F2937")
OUTCOME_CHANGE_FILL = PatternFill(fill_type="solid", fgColor="FDE68A")
PROXY_CHANGE_FILL = PatternFill(fill_type="solid", fgColor="93C5FD")
NOTE_FILL = PatternFill(fill_type="solid", fgColor="BBF7D0")
ORIGINAL_PROTECTED_COLUMN_FILL = PatternFill(fill_type="solid", fgColor="FEE2E2")
ORIGINAL_PROXY_COLUMN_FILL = PatternFill(fill_type="solid", fgColor="E0F2FE")
SUGGESTED_OUTCOME_COLUMN_FILL = PatternFill(fill_type="solid", fgColor="FEF3C7")
SUGGESTED_PROXY_COLUMN_FILL = PatternFill(fill_type="solid", fgColor="D1FAE5")
RISK_FILLS = {
    "HIGH RISK": PatternFill(fill_type="solid", fgColor="FCA5A5"),
    "MEDIUM RISK": PatternFill(fill_type="solid", fgColor="FCD34D"),
    "LOW RISK": PatternFill(fill_type="solid", fgColor="86EFAC"),
}
POSITIVE_OUTCOMES = {
    "approved", "approve", "accepted", "accept", "yes", "y", "true", "1", "hired", "loan_granted",
    "granted", "selected", "high", "pass", "passed", "success",
}
INVALID_TEXT_VALUES = {"", "missing", "unknown", "nan", "none", "n/a", "na", "generalized"}
TEXT_VALUE_POOLS = {
    "surname": [
        "Smith", "Johnson", "Williams", "Brown", "Jones", "Garcia", "Miller", "Davis", "Rodriguez",
        "Martinez", "Hernandez", "Lopez", "Gonzalez", "Wilson", "Anderson", "Thomas", "Taylor",
        "Moore", "Jackson", "Martin", "Lee", "Perez", "Thompson", "White", "Harris", "Sanchez",
        "Clark", "Ramirez", "Lewis", "Robinson", "Walker", "Young", "Allen", "King", "Wright",
        "Scott", "Torres", "Nguyen", "Hill", "Flores", "Green", "Adams", "Nelson", "Baker", "Hall",
        "Rivera", "Campbell", "Mitchell", "Carter", "Roberts",
    ],
    "location": [
        "Downtown", "Riverside", "Hillcrest", "Lakeside", "Northgate", "Westside", "Eastwood",
        "Oak Park", "Midtown", "Harbor View", "Maple Heights", "Cedar Grove", "Brookfield",
        "Sunnyvale", "Fairview", "Old Town", "Greenfield", "River Park", "Pine Ridge", "West End",
        "Elmwood", "Meadowbrook", "Silver Lake", "Highland Park",
    ],
    "institution": [
        "North Ridge University", "Metro State College", "Lakeshore Institute", "Central City University",
        "Horizon College", "Riverview University", "Summit Institute of Technology", "Valley State College",
        "Pioneer University", "Crescent Academy", "Harbor University", "Prairie College",
    ],
    "language": [
        "English", "Spanish", "Hindi", "Mandarin", "French", "Arabic", "Portuguese", "Bengali",
        "German", "Japanese", "Urdu", "Russian", "Korean", "Turkish", "Vietnamese",
    ],
    "occupation": [
        "Teacher", "Nurse", "Engineer", "Analyst", "Technician", "Supervisor", "Sales Associate",
        "Driver", "Clerk", "Manager", "Consultant", "Operator", "Coordinator", "Developer",
        "Assistant", "Specialist", "Accountant", "Designer", "Electrician", "Carpenter",
    ],
    "status": [
        "Full-time", "Part-time", "Contract", "Self-employed", "Unemployed", "Retired", "Student",
    ],
    "insurance": [
        "Basic Plan", "Standard Plan", "Premium Plan", "Employer-Sponsored", "Marketplace Plan",
        "Family Plan", "Individual Plan", "HMO", "PPO", "EPO", "Medicaid", "Medicare",
    ],
    "organization": [
        "Community Association", "Sports Club", "Alumni Association", "Volunteer Group",
        "Neighborhood Council", "Civic Society", "Cultural Club", "Professional Guild",
        "Arts Collective", "Trade Association",
    ],
}
GENERIC_TEXT_POOL = [
    "Riverside", "Hillcrest", "Maple Grove", "Brookside", "Cedar Park", "Fairview", "Lakeside",
    "Oak Ridge", "Pine Hollow", "Sunset Park", "Meadow View", "Highland Grove",
]
_LLM_PROXY_POOL_CACHE: Dict[str, List[str]] = {}


def _safe_stem(filename: str) -> str:
    stem = Path(filename or "dataset").stem
    cleaned = re.sub(r"[^A-Za-z0-9_-]+", "_", stem).strip("_")
    return cleaned or "dataset"


def _next_report_path(filename: str) -> Path:
    stem = _safe_stem(filename)
    for index in range(1, 1000):
        candidate = REPORT_DIR / f"{stem}{index:03d}.xlsx"
        if not candidate.exists():
            return candidate
    return REPORT_DIR / f"{stem}{math.floor(Path.cwd().stat().st_mtime)}.xlsx"


def _load_dataframe(dataset_path: str):
    if pd is None:
        raise RuntimeError("pandas is required to generate the Excel report.")
    path = Path(dataset_path)
    suffix = path.suffix.lower()
    if suffix == ".csv":
        frame = pd.read_csv(path, header=0)
    elif suffix in {".xlsx", ".xlsm"}:
        sheet_name = select_xlsx_sheet_name(path)
        header_row = find_header_row_index(path, sheet_name)
        frame = pd.read_excel(path, sheet_name=sheet_name, header=header_row)
    elif suffix == ".xls":
        frame = pd.read_excel(path, header=0, engine="xlrd")
    else:
        raise RuntimeError(f"Unsupported file type: {suffix}")
    return _clean_dataframe(frame).reset_index(drop=True)


def _normalize_text(value: Any) -> str:
    return str(value).strip().lower()


def _is_positive(value: Any, positive_label: str) -> bool:
    text = _normalize_text(value)
    return text == _normalize_text(positive_label) or text in POSITIVE_OUTCOMES


def _detect_positive_label(frame, outcome_column: str, analysis: Dict[str, Any]) -> str:
    distribution = ((analysis.get("outcome_column") or {}).get("distribution") or {})
    preferred = ["approved", "yes", "hired", "accepted", "selected", "granted", "high", "true", "1"]
    lowered = {str(key).strip().lower(): key for key in distribution.keys()}
    for label in preferred:
        if label in lowered:
            return str(lowered[label])
    if outcome_column in frame.columns:
        values = [_normalize_text(value) for value in frame[outcome_column].dropna().tolist()]
        for label in preferred:
            if label in values:
                return label
    return "approved"


def _coerce_like_series(series, value: Any) -> Any:
    if pd is None:
        return value
    try:
        if pd.api.types.is_integer_dtype(series.dtype):
            return int(float(str(value)))
        if pd.api.types.is_float_dtype(series.dtype):
            return float(str(value))
    except Exception:
        return value
    return value


def _normalize_list_values(values: Any) -> List[str]:
    if not isinstance(values, list):
        return []
    cleaned: List[str] = []
    seen: set[str] = set()
    for value in values:
        text = str(value).strip()
        normalized = _normalize_text(text)
        if not text or normalized in INVALID_TEXT_VALUES:
            continue
        if normalized in seen:
            continue
        seen.add(normalized)
        cleaned.append(text)
    return cleaned


def _age_bucket(value: Any) -> str:
    try:
        number = float(str(value).replace(",", "").strip())
    except Exception:
        return "unknown"
    if not math.isfinite(number):
        return "unknown"
    if number < 25:
        return "Under 25"
    if number <= 34:
        return "25 - 34"
    if number <= 44:
        return "35 - 44"
    if number <= 54:
        return "45 - 54"
    if number <= 64:
        return "55 - 64"
    return "65 and above"


def _matches_group(column_name: str, raw_value: Any, group_value: str) -> bool:
    lowered = str(column_name).lower()
    if any(token in lowered for token in ["age", "dob", "birth"]):
        return _age_bucket(raw_value) == str(group_value)
    return _normalize_text(raw_value) == _normalize_text(group_value)


def _observed_text_pool(frame, column_name: str) -> List[str]:
    if pd is None or frame is None or column_name not in frame.columns:
        return []
    values: List[str] = []
    seen: set[str] = set()
    for raw in frame[column_name].tolist():
        text = str(raw).strip()
        normalized = _normalize_text(text)
        if not text or normalized in INVALID_TEXT_VALUES or "*" in text:
            continue
        if normalized in seen:
            continue
        seen.add(normalized)
        values.append(text)
    return values


def _preferred_proxy_pool(column_name: str) -> List[str]:
    lowered = str(column_name).lower()
    if any(token in lowered for token in ["surname", "last_name", "lastname"]):
        return TEXT_VALUE_POOLS["surname"]
    if any(token in lowered for token in ["district", "area", "neighborhood", "city", "town"]):
        return TEXT_VALUE_POOLS["location"]
    if any(token in lowered for token in ["school", "university", "college", "institute"]):
        return TEXT_VALUE_POOLS["institution"]
    if any(token in lowered for token in ["language", "mother_tongue"]):
        return TEXT_VALUE_POOLS["language"]
    if any(token in lowered for token in ["occupation", "job_type"]):
        return TEXT_VALUE_POOLS["occupation"]
    if "employment_status" in lowered:
        return TEXT_VALUE_POOLS["status"]
    if "insurance_type" in lowered:
        return TEXT_VALUE_POOLS["insurance"]
    if any(token in lowered for token in ["club", "association"]):
        return TEXT_VALUE_POOLS["organization"]
    return GENERIC_TEXT_POOL


def _pick_pool_value(pool: List[str], column_name: str, row_index: int, current_value: Any) -> str:
    normalized_current = _normalize_text(current_value)
    candidates = [item for item in pool if _normalize_text(item) not in INVALID_TEXT_VALUES and _normalize_text(item) != normalized_current]
    if not candidates:
        candidates = [item for item in pool if _normalize_text(item) not in INVALID_TEXT_VALUES]
    if not candidates:
        return str(current_value).strip() or "Suggested"
    rng = random.Random(f"{column_name}|{row_index}|{current_value}")
    choice = candidates[rng.randrange(len(candidates))]
    if _normalize_text(choice) == normalized_current and len(candidates) > 1:
        choice = candidates[(rng.randrange(len(candidates)) + 1) % len(candidates)]
    return choice


def _suggest_zip_code(column_name: str, value: Any, row_index: int, frame) -> str:
    text = str(value).strip()
    normalized_current = re.sub(r"\D", "", text)

    observed_zips = []
    if pd is not None and frame is not None and column_name in frame.columns:
        for raw in frame[column_name].tolist():
            digits = re.sub(r"\D", "", str(raw))
            if len(digits) == 5:
                observed_zips.append(digits)
        observed_zips = list(dict.fromkeys(observed_zips))
    if len(observed_zips) > 1:
        candidates = [item for item in observed_zips if item != normalized_current]
        if not candidates:
            candidates = observed_zips
        rng = random.Random(f"{column_name}|{row_index}|{text}|zip")
        choice = candidates[rng.randrange(len(candidates))]
        if choice == normalized_current and len(candidates) > 1:
            choice = candidates[(rng.randrange(len(candidates)) + 1) % len(candidates)]
        return choice

    prefix = normalized_current[:3] if len(normalized_current) >= 3 else ""
    if not prefix and pd is not None and frame is not None and column_name in frame.columns:
        for raw in frame[column_name].tolist():
            digits = re.sub(r"\D", "", str(raw))
            if len(digits) >= 3:
                prefix = digits[:3]
                break
    if not prefix:
        prefixes = ["100", "112", "200", "303", "606", "752", "770", "850", "900"]
        rng = random.Random(f"{column_name}|{row_index}|{text}|prefix")
        prefix = prefixes[rng.randrange(len(prefixes))]

    rng = random.Random(f"{column_name}|{row_index}|{text}|suffix")
    suffix = rng.randrange(100)
    candidate = f"{prefix}{suffix:02d}"
    if len(normalized_current) >= 5 and candidate == normalized_current[:5]:
        candidate = f"{prefix}{(suffix + 17) % 100:02d}"
    return candidate


def _suggest_proxy_value(column_name: str, value: Any, row_index: int, frame) -> Any:
    if value is None or str(value).strip() == "":
        return value
    lowered = str(column_name).lower()
    if any(token in lowered for token in ["zip", "postal", "postcode"]):
        return _suggest_zip_code(column_name, value, row_index, frame)

    observed_pool = _observed_text_pool(frame, column_name)
    pool = observed_pool if len(observed_pool) > 1 else _preferred_proxy_pool(column_name)
    return _pick_pool_value(pool, column_name, row_index, value)


def _proxy_pool_cache_key(column_name: str, display_name: str, observed_values: List[str], reason: str) -> str:
    sample = "|".join(observed_values[:8])
    return f"{column_name}::{display_name}::{reason}::{sample}"


def _llm_proxy_pool(
    column_name: str,
    display_name: str,
    observed_values: List[str],
    finding: Dict[str, Any],
) -> List[str]:
    reason = str(finding.get("reason") or "")
    cache_key = _proxy_pool_cache_key(column_name, display_name, observed_values, reason)
    cached = _LLM_PROXY_POOL_CACHE.get(cache_key)
    if cached:
        return cached

    payload = {
        "column_name": column_name,
        "display_name": display_name,
        "reason": reason,
        "observed_values": observed_values[:8],
        "risk": finding.get("risk"),
        "finding_type": finding.get("type"),
    }
    prompt = (
        "You are helping create a realistic suggested-data Excel sheet for a fairness report. "
        "Return only JSON with a key named suggested_values containing a list of 12 realistic replacement values "
        "for the flagged column. Keep the values the same general type as the column, varied, and natural. "
        "Avoid placeholders like Generalized, Unknown, Sample, or Test. "
        "If the column is a name/surname, return common surnames. If it is a location column, return real place-like values. "
        "If it is an organization, school, occupation, insurance, or similar column, return realistic matching values. "
        "Use these details: "
        f"{json.dumps(payload, ensure_ascii=False)}"
    )

    try:
        response = complete_json(prompt, fallback={})
    except Exception:
        response = {}

    suggestions = _normalize_list_values(response.get("suggested_values")) if isinstance(response, dict) else []
    if not suggestions:
        suggestions = []

    combined: List[str] = []
    seen: set[str] = set()
    for candidate in [*suggestions, *observed_values, *_preferred_proxy_pool(column_name)]:
        text = str(candidate).strip()
        normalized = _normalize_text(text)
        if not text or normalized in INVALID_TEXT_VALUES or normalized in seen:
            continue
        seen.add(normalized)
        combined.append(text)

    if not combined:
        combined = _preferred_proxy_pool(column_name)

    _LLM_PROXY_POOL_CACHE[cache_key] = combined
    return combined


def _findings_column_map(analysis: Dict[str, Any]) -> Dict[str, str]:
    mapping: Dict[str, str] = {}
    for finding in analysis.get("bias_findings") or []:
        column = finding.get("column")
        if column:
            mapping[str(column)] = "protected"
    for finding in analysis.get("proxy_findings") or []:
        column = finding.get("column")
        if column:
            mapping[str(column)] = "proxy"
    outcome_column = (analysis.get("outcome_column") or {}).get("original_column") or (analysis.get("outcome_column") or {}).get("column")
    if outcome_column:
        mapping[str(outcome_column)] = "outcome"
    return mapping


def _column_highlights_from_analysis(analysis: Dict[str, Any]) -> Dict[str, PatternFill]:
    highlights: Dict[str, PatternFill] = {}
    for finding in analysis.get("bias_findings") or []:
        column = finding.get("column")
        if column:
            highlights[str(column)] = ORIGINAL_PROTECTED_COLUMN_FILL
    for finding in analysis.get("proxy_findings") or []:
        column = finding.get("column")
        if column:
            highlights[str(column)] = ORIGINAL_PROXY_COLUMN_FILL
    return highlights


def _column_highlights_from_changes(changes: List[Dict[str, Any]]) -> Dict[str, PatternFill]:
    highlights: Dict[str, PatternFill] = {}
    for change in changes:
        column = str(change.get("column") or "")
        if not column:
            continue
        highlights[column] = SUGGESTED_OUTCOME_COLUMN_FILL if change.get("kind") == "outcome" else SUGGESTED_PROXY_COLUMN_FILL
    return highlights


def _apply_column_highlight(worksheet, frame, highlights: Dict[str, PatternFill]) -> None:
    if not highlights:
        return
    column_lookup = {str(name): position + 1 for position, name in enumerate(frame.columns)}
    for column_name, fill in highlights.items():
        column_number = column_lookup.get(str(column_name))
        if not column_number:
            continue
        for row_number in range(2, worksheet.max_row + 1):
            worksheet.cell(row=row_number, column=column_number).fill = fill


def _merge_value_pools(*pools: List[str]) -> List[str]:
    merged: List[str] = []
    seen: set[str] = set()
    for pool in pools:
        for value in pool or []:
            text = str(value).strip()
            normalized = _normalize_text(text)
            if not text or normalized in INVALID_TEXT_VALUES or normalized in seen:
                continue
            seen.add(normalized)
            merged.append(text)
    return merged


def _proxy_value_pool_for_finding(frame, finding: Dict[str, Any]) -> List[str]:
    column_name = str(finding.get("column") or "")
    display_name = str(finding.get("display_name") or column_name or "Column")
    observed_pool = _observed_text_pool(frame, column_name)
    llm_pool = _llm_proxy_pool(column_name, display_name, observed_pool, finding)
    merged_pool = _merge_value_pools(llm_pool, observed_pool, _preferred_proxy_pool(column_name))
    return merged_pool or _preferred_proxy_pool(column_name)


def _protected_lookup(analysis: Dict[str, Any]) -> Dict[str, Dict[str, Any]]:
    items = analysis.get("protected_columns") or []
    return {
        str(item.get("original_column") or item.get("column")): item
        for item in items
        if isinstance(item, dict) and (item.get("original_column") or item.get("column"))
    }


def _build_suggested_dataframe(frame, analysis: Dict[str, Any]) -> Tuple[Any, List[Dict[str, Any]]]:
    adjusted = frame.copy(deep=True)
    notes: Dict[int, List[str]] = {}
    changes: List[Dict[str, Any]] = []
    proxy_pools: Dict[str, List[str]] = {}

    outcome_column = (analysis.get("outcome_column") or {}).get("original_column") or (analysis.get("outcome_column") or {}).get("column")
    positive_label = _detect_positive_label(adjusted, outcome_column, analysis) if outcome_column in adjusted.columns else "approved"
    positive_value = _coerce_like_series(adjusted[outcome_column], positive_label) if outcome_column in adjusted.columns else positive_label
    protected_meta = _protected_lookup(analysis)

    for finding in analysis.get("bias_findings") or []:
        column = finding.get("column")
        if column not in adjusted.columns or outcome_column not in adjusted.columns:
            continue
        matching_rows = adjusted.index[adjusted[column].apply(lambda value: _matches_group(column, value, finding.get("biased_value")))].tolist()
        if not matching_rows:
            continue
        eligible_rows = [idx for idx in matching_rows if not _is_positive(adjusted.at[idx, outcome_column], positive_label)]
        if not eligible_rows:
            continue

        group_size = len(matching_rows)
        current_positive = sum(1 for idx in matching_rows if _is_positive(adjusted.at[idx, outcome_column], positive_label))
        target_rate = max(float(finding.get("biased_rate") or 0), float(finding.get("reference_rate") or 0) - 4.0)
        target_positive = math.ceil((target_rate / 100.0) * group_size)
        flips_needed = min(max(0, target_positive - current_positive), len(eligible_rows))

        for idx in eligible_rows[:flips_needed]:
            old_value = adjusted.at[idx, outcome_column]
            adjusted.at[idx, outcome_column] = positive_value
            notes.setdefault(idx, []).append(
                f"Adjusted {outcome_column} from {old_value} to {positive_value} for {column} = {finding.get('biased_value')}"
            )
            changes.append(
                {
                    "row_index": idx,
                    "column": outcome_column,
                    "old_value": old_value,
                    "new_value": positive_value,
                    "kind": "outcome",
                }
            )

        meta = protected_meta.get(str(column), {})
        if meta.get("population_share") is not None:
            notes.setdefault(matching_rows[0], []).append(
                f"Review {column}: {finding.get('biased_value')} was flagged with a {finding.get('gap')}% approval gap."
            )

    for finding in analysis.get("proxy_findings") or []:
        column = finding.get("column")
        if column not in adjusted.columns:
            continue
        adjusted[column] = adjusted[column].astype(object)
        pool = proxy_pools.get(str(column))
        if pool is None:
            pool = _proxy_value_pool_for_finding(frame, finding)
            proxy_pools[str(column)] = pool
        for idx in adjusted.index.tolist():
            old_value = adjusted.at[idx, column]
            if any(token in str(column).lower() for token in ["zip", "postal", "postcode"]):
                new_value = _suggest_zip_code(column, old_value, idx, frame)
            else:
                new_value = _pick_pool_value(pool, column, idx, old_value)
            if _normalize_text(old_value) == _normalize_text(new_value):
                continue
            adjusted.at[idx, column] = new_value
            notes.setdefault(idx, []).append(f"Suggested {column} from {old_value} to {new_value}")
            changes.append(
                {
                    "row_index": idx,
                    "column": column,
                    "old_value": old_value,
                    "new_value": new_value,
                    "kind": "proxy",
                }
            )

    return adjusted, changes


def _style_headers(worksheet) -> None:
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL


def _autosize(worksheet, max_width: int = 42) -> None:
    for column_cells in worksheet.columns:
        length = 0
        letter = get_column_letter(column_cells[0].column)
        for cell in column_cells[:120]:
            try:
                length = max(length, len(str(cell.value or "")))
            except Exception:
                continue
        worksheet.column_dimensions[letter].width = min(max(length + 2, 12), max_width)


def _write_dataframe_sheet(
    workbook: Workbook,
    title: str,
    frame,
    changes: List[Dict[str, Any]] | None = None,
    column_highlights: Dict[str, PatternFill] | None = None,
) -> None:
    worksheet = workbook.create_sheet(title)
    for row in dataframe_to_rows(frame, index=False, header=True):
        worksheet.append(row)
    _style_headers(worksheet)

    if column_highlights:
        _apply_column_highlight(worksheet, frame, column_highlights)

    if changes:
        column_lookup = {str(name): position + 1 for position, name in enumerate(frame.columns)}
        for change in changes:
            row_number = int(change["row_index"]) + 2
            column_number = column_lookup.get(str(change["column"]))
            if not column_number:
                continue
            worksheet.cell(row=row_number, column=column_number).fill = (
                OUTCOME_CHANGE_FILL if change.get("kind") == "outcome" else PROXY_CHANGE_FILL
            )

    worksheet.freeze_panes = "A2"
    _autosize(worksheet)


def _add_report_sheet(workbook: Workbook, analysis: Dict[str, Any], source_name: str) -> None:
    worksheet = workbook.active
    worksheet.title = "Bias_Report"

    worksheet["A1"] = "FairTrace Dataset Bias Report"
    worksheet["A1"].font = Font(bold=True, size=16)
    worksheet["A3"] = "Source Dataset"
    worksheet["B3"] = source_name
    worksheet["A4"] = "Overall Risk"
    worksheet["B4"] = analysis.get("overall_risk") or "LOW RISK"
    worksheet["B4"].fill = RISK_FILLS.get(str(worksheet["B4"].value), RISK_FILLS["LOW RISK"])
    worksheet["A5"] = "Summary"
    worksheet["B5"] = analysis.get("dataset_summary") or "No summary available."
    worksheet["A6"] = "Rows Checked"
    worksheet["B6"] = analysis.get("row_count") or 0

    protected_header_row = 9
    worksheet[f"A{protected_header_row}"] = "Biased Columns"
    worksheet[f"A{protected_header_row}"].font = Font(bold=True)
    protected_table_row = protected_header_row + 1
    protected_headers = ["Column", "Biased Value", "Biased Rate", "Reference Value", "Reference Rate", "Gap", "Risk", "Recommendation"]
    for offset, heading in enumerate(protected_headers, start=1):
        cell = worksheet.cell(row=protected_table_row, column=offset, value=heading)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL

    for row_offset, item in enumerate(analysis.get("bias_findings") or [], start=1):
        values = [
            item.get("display_name"),
            item.get("biased_value"),
            item.get("biased_rate"),
            item.get("reference_value"),
            item.get("reference_rate"),
            item.get("gap"),
            item.get("risk"),
            item.get("recommendation"),
        ]
        for col_offset, value in enumerate(values, start=1):
            worksheet.cell(row=protected_table_row + row_offset, column=col_offset, value=value)

    proxy_section_row = max(protected_table_row + max(len(analysis.get("bias_findings") or []), 1) + 3, 16)
    worksheet[f"A{proxy_section_row}"] = "Proxy Warnings"
    worksheet[f"A{proxy_section_row}"].font = Font(bold=True)
    proxy_header_row = proxy_section_row + 1
    proxy_headers = ["Column", "Flagged Value", "Share", "Risk", "Recommendation"]
    for offset, heading in enumerate(proxy_headers, start=1):
        cell = worksheet.cell(row=proxy_header_row, column=offset, value=heading)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL

    for row_offset, item in enumerate(analysis.get("proxy_findings") or [], start=1):
        values = [
            item.get("display_name"),
            item.get("biased_value"),
            item.get("biased_share"),
            item.get("risk"),
            item.get("recommendation"),
        ]
        for col_offset, value in enumerate(values, start=1):
            worksheet.cell(row=proxy_header_row + row_offset, column=col_offset, value=value)

    _autosize(worksheet)


def _chart_asset_path(original_filename: str, kind: str) -> Path:
    prefix = f"{_safe_stem(original_filename)}_{kind}_"
    fd, raw_path = tempfile.mkstemp(prefix=prefix, suffix=".png", dir=REPORT_DIR)
    os.close(fd)
    return Path(raw_path)


def _render_bar_chart_image(labels: List[str], values: List[float], title: str, xlabel: str, ylabel: str, output_path: Path) -> None:
    if plt is None:
        raise RuntimeError("matplotlib is required to render chart images.")
    figure, axis = plt.subplots(figsize=(8.6, 4.6), dpi=140)
    palette = ["#2563EB", "#0EA5E9", "#10B981", "#F59E0B", "#F97316", "#8B5CF6", "#14B8A6", "#EF4444"]
    bars = axis.bar(labels, values, color=[palette[index % len(palette)] for index in range(len(labels))], edgecolor="#1F2937", linewidth=0.6)
    axis.set_title(title, fontsize=13, fontweight="bold", pad=12)
    axis.set_xlabel(xlabel, fontsize=10)
    axis.set_ylabel(ylabel, fontsize=10)
    axis.grid(axis="y", alpha=0.25, linestyle="--")
    axis.set_axisbelow(True)
    axis.spines["top"].set_visible(False)
    axis.spines["right"].set_visible(False)
    if labels:
        axis.tick_params(axis="x", labelrotation=20)
    if values:
        axis.set_ylim(0, max(values) * 1.25 if max(values) > 0 else 1)
    for bar, value in zip(bars, values):
        text = f"{value:.1f}%" if any(token in ylabel.lower() for token in ("rate", "share", "percent")) else f"{value:g}"
        axis.annotate(
            text,
            (bar.get_x() + bar.get_width() / 2.0, bar.get_height()),
            xytext=(0, 4),
            textcoords="offset points",
            ha="center",
            va="bottom",
            fontsize=8,
            color="#111827",
        )
    figure.tight_layout()
    figure.savefig(output_path, bbox_inches="tight", facecolor="white")
    plt.close(figure)


def _render_pie_chart_image(labels: List[str], values: List[float], title: str, output_path: Path) -> None:
    if plt is None:
        raise RuntimeError("matplotlib is required to render chart images.")
    figure, axis = plt.subplots(figsize=(8.0, 4.8), dpi=140)
    palette = ["#2563EB", "#F59E0B", "#10B981", "#EF4444", "#8B5CF6", "#0EA5E9"]
    total = sum(values)
    if total <= 0:
        labels = ["No findings"]
        values = [1]
        colors = ["#CBD5E1"]
    else:
        colors = [palette[index % len(palette)] for index in range(len(values))]
    wedges, texts, autotexts = axis.pie(
        values,
        labels=labels,
        autopct=(lambda pct: f"{pct:.1f}%" if total > 0 else ""),
        startangle=90,
        colors=colors,
        textprops={"fontsize": 9},
    )
    for autotext in autotexts:
        autotext.set_color("white")
        autotext.set_fontweight("bold")
    axis.set_title(title, fontsize=13, fontweight="bold", pad=12)
    axis.axis("equal")
    figure.tight_layout()
    figure.savefig(output_path, bbox_inches="tight", facecolor="white")
    plt.close(figure)


def _add_chart_image(worksheet, image_path: Path, anchor: str, width: int, height: int) -> None:
    image = XLImage(str(image_path))
    image.width = width
    image.height = height
    worksheet.add_image(image, anchor)


def _add_report_charts(worksheet, analysis: Dict[str, Any], original_filename: str = "report") -> List[Path]:
    if plt is None or XLImage is None:
        raise RuntimeError(
            "matplotlib and Pillow are required to render workbook charts. "
            "Install backend/requirements.txt before generating the Excel report."
        )

    bias_findings = analysis.get("bias_findings") or []
    proxy_findings = analysis.get("proxy_findings") or []
    recommendation_count = len(analysis.get("overall_recommendations") or [])

    protected_columns = {
        str(item.get("original_column") or item.get("column")): item
        for item in analysis.get("protected_columns") or []
        if isinstance(item, dict)
    }

    table_end_row = max(
        9 + 1 + max(len(bias_findings), 1),
        16 + 1 + max(len(proxy_findings), 1),
    )
    chart_start_row = max(table_end_row + 4, 28)
    worksheet.cell(row=chart_start_row, column=1, value="Charts")
    worksheet.cell(row=chart_start_row, column=1).font = Font(bold=True)

    bar_anchor_row = chart_start_row + 1
    pie_anchor_row = bar_anchor_row + 20
    created_assets: List[Path] = []

    first_finding = (bias_findings or [None])[0]
    if first_finding:
        meta = protected_columns.get(str(first_finding.get("column")))
        if meta and (meta.get("distribution") or {}):
            distribution = dict(meta.get("distribution") or {})
            labels = list(distribution.keys())[:8]
            values = [float(distribution[label]) for label in labels]
            bar_path = _chart_asset_path(original_filename, "bar")
            _render_bar_chart_image(
                labels=labels,
                values=values,
                title=f"{first_finding.get('display_name')} Approval Rates",
                xlabel="Group",
                ylabel="Approval Rate (%)",
                output_path=bar_path,
            )
            _add_chart_image(worksheet, bar_path, f"J{bar_anchor_row}", 760, 360)
            created_assets.append(bar_path)

    if not created_assets and proxy_findings:
        labels = [str(item.get("display_name") or item.get("column") or "Proxy") for item in proxy_findings[:8]]
        values = [float(item.get("biased_share") or 0) for item in proxy_findings[:8]]
        bar_path = _chart_asset_path(original_filename, "bar")
        _render_bar_chart_image(
            labels=labels,
            values=values,
            title="Proxy Concentration",
            xlabel="Column",
            ylabel="Dominant Share (%)",
            output_path=bar_path,
        )
        _add_chart_image(worksheet, bar_path, f"J{bar_anchor_row}", 760, 360)
        created_assets.append(bar_path)

    if not created_assets:
        labels = ["Biased Columns", "Proxy Warnings"]
        values = [float(len(bias_findings)), float(len(proxy_findings))]
        bar_path = _chart_asset_path(original_filename, "bar")
        _render_bar_chart_image(
            labels=labels,
            values=values,
            title="Finding Summary",
            xlabel="Finding",
            ylabel="Count",
            output_path=bar_path,
        )
        _add_chart_image(worksheet, bar_path, f"J{bar_anchor_row}", 760, 360)
        created_assets.append(bar_path)

    pie_labels = ["Biased Columns", "Proxy Warnings"]
    pie_values = [float(len(bias_findings)), float(len(proxy_findings))]
    pie_path = _chart_asset_path(original_filename, "pie")
    _render_pie_chart_image(
        labels=pie_labels,
        values=pie_values,
        title="Finding Mix",
        output_path=pie_path,
    )
    _add_chart_image(worksheet, pie_path, f"J{pie_anchor_row}", 720, 380)
    created_assets.append(pie_path)

    recommendation_row = pie_anchor_row + 22
    worksheet[f"A{recommendation_row}"] = "Recommended Changes"
    worksheet[f"A{recommendation_row}"].font = Font(bold=True)
    for index, text in enumerate(analysis.get("overall_recommendations") or [], start=1):
        worksheet.cell(row=recommendation_row + index, column=1, value=text)

    _autosize(worksheet)
    return created_assets


def create_analysis_workbook(dataset_path: str, original_filename: str, analysis: Dict[str, Any]) -> Dict[str, str]:
    frame = _load_dataframe(dataset_path)
    suggested_frame, changes = _build_suggested_dataframe(frame, analysis)

    workbook = Workbook()
    _add_report_sheet(workbook, analysis, original_filename)
    chart_assets = _add_report_charts(workbook["Bias_Report"], analysis, original_filename)
    _write_dataframe_sheet(workbook, "Original_Data", frame, column_highlights=_column_highlights_from_analysis(analysis))
    _write_dataframe_sheet(workbook, "Updated_Data", suggested_frame, changes)

    target_path = _next_report_path(original_filename)
    try:
        workbook.save(target_path)
    finally:
        for asset_path in chart_assets:
            try:
                Path(asset_path).unlink(missing_ok=True)
            except Exception:
                pass
    return {
        "filename": target_path.name,
        "path": str(target_path),
    }
